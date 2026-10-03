"""Offline tests for the real-data / applicant-document layer.

Network calls are intentionally not exercised here: external public boards are
contracted at runtime and the app reports provenance / source failure rather
than fabricating their data.
"""
from __future__ import annotations

import asyncio
import copy
import importlib
from concurrent.futures import ThreadPoolExecutor
import sys
import tempfile
import unittest
import warnings
from unittest.mock import AsyncMock, patch
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
warnings.filterwarnings("ignore", message=r"Using `httpx` with `starlette\.testclient` is deprecated")

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader
from reportlab.pdfgen.canvas import Canvas
import httpx

from app.data import DEMO_CANDIDATE, DEMO_JOBS
from app.documents import _application_bullets, build_docx, build_pdf, export_readiness
from app.evidence import build_cover_letter, build_safe_patches, cover_letter_evidence_ids, extract_evidence_from_resume, score_job, validate_patches
from app.main import _candidate_from_text, app, candidate_as_dict, execute_evidence_lab
from app.schemas import CandidateInput, RunStatus
main_module = importlib.import_module("app.main")
from app.sources import MAX_PUBLIC_RESPONSE_BYTES, clean_text, display_posted, fetch_greenhouse, fetch_public_json, greenhouse_token, lever_slug, safe_http_url


class RealWorkflowTests(unittest.TestCase):
    def _result(self):
        candidate = copy.deepcopy(DEMO_CANDIDATE)
        job = copy.deepcopy(DEMO_JOBS[0])
        match = score_job(job, candidate)
        patches = validate_patches(build_safe_patches(job, candidate, match), candidate)["safe_patches"]
        return {
            "candidate": {key: candidate[key] for key in ("name", "headline", "location")},
            "candidate_full": candidate,
            "job": job,
            "match": match,
            "patches": patches,
            "cover_letter": build_cover_letter(job, candidate, match),
        }

    def test_demo_candidate_identity_and_schema_default_are_hussain(self):
        self.assertEqual(DEMO_CANDIDATE["name"], "Hussain Ahmed")
        self.assertEqual(DEMO_CANDIDATE["headline"], "Full-Stack Developer · Python · FastAPI · React · PostgreSQL")
        self.assertEqual(DEMO_CANDIDATE["location"], "Rawalpindi, Pakistan")
        self.assertEqual(CandidateInput().name, "Hussain Ahmed")
        self.assertEqual(CandidateInput().headline, "Full-Stack Developer · Python · FastAPI · React · PostgreSQL")

    def test_cv_parser_keeps_common_explicit_pakistan_locations(self):
        for line, expected in [
            ("Rawalpindi, Pakistan", "Rawalpindi, Pakistan"),
            ("Islamabad, Pakistan", "Islamabad, Pakistan"),
            ("Lahore, Pakistan", "Lahore, Pakistan"),
            ("Location: Rawalpindi", "Rawalpindi"),
            ("Based in: Rawalpindi", "Rawalpindi"),
        ]:
            with self.subTest(line=line):
                candidate = _candidate_from_text(f"Hussain Ahmed\nFull-Stack Developer\n{line}\nEXPERIENCE\nBuilt and containerized a Docker service for a student platform used by project collaborators.\nSKILLS\nDocker", "hussain.txt")
                self.assertEqual(candidate.location, expected)

    def test_hussain_style_cv_regression_keeps_location_metrics_and_docker_proof(self):
        resume = """Hussain Ahmed
Full-Stack Developer
Rawalpindi, Pakistan
EXPERIENCE
• Containerized a FastAPI service with Docker and PostgreSQL for 2,000 users.
• Improved reporting by 40 percent and supported $20,000 in subscriptions over 3.5 years.
SKILLS
Python, FastAPI, PostgreSQL, Docker
"""
        candidate = _candidate_from_text(resume, "hussain-ahmed.txt")
        self.assertEqual(candidate.location, "Rawalpindi, Pakistan")
        metrics = "\n".join(item.metric or "" for item in candidate.evidence)
        for expected in ("2,000 users", "40 percent", "$20,000", "3.5 years"):
            self.assertIn(expected, metrics)
        match = score_job({"must_have": ["Docker", "FastAPI", "PostgreSQL"], "preferred": []}, candidate.model_dump())
        self.assertIn("Docker", {item["skill"] for item in match["verified_strengths"]})
        self.assertNotIn("Docker", {item["skill"] for item in match["gaps"]})
        self.assertEqual(extract_evidence_from_resume("SKILLS: Docker")[0]["skills"], ["Docker"])
        response = TestClient(app).post("/api/candidate/upload", files={"file": ("hussain-ahmed.txt", resume.encode("utf-8"), "text/plain")})
        self.assertEqual(response.status_code, 200, response.text)
        uploaded = response.json()["candidate"]
        self.assertEqual(uploaded["location"], "Rawalpindi, Pakistan")
        self.assertIn("Docker", {skill for item in uploaded["evidence"] for skill in item["skills"]})

    def test_public_board_urls_extract_only_the_public_board_identifier(self):
        self.assertEqual(greenhouse_token("https://boards.greenhouse.io/stripe/jobs/123"), "stripe")
        self.assertEqual(greenhouse_token("stripe"), "stripe")
        self.assertEqual(lever_slug("https://jobs.lever.co/metabase/abc-def"), "metabase")
        self.assertEqual(lever_slug("metabase"), "metabase")

    def test_public_board_urls_reject_an_unrelated_host(self):
        with self.assertRaises(ValueError):
            greenhouse_token("https://evil.example/stripe")
        with self.assertRaises(ValueError):
            lever_slug("https://evil.example/metabase")

    def test_board_identifiers_require_https_allowlisted_urls_but_keep_bare_tokens_safe(self):
        self.assertEqual(greenhouse_token("stripe"), "stripe")
        self.assertEqual(lever_slug("metabase"), "metabase")
        for unsafe in (
            "http://boards.greenhouse.io/stripe",
            "https://user:pass@boards.greenhouse.io/stripe",
            "https://localhost/stripe",
            "https://127.0.0.1/stripe",
            "https://boards.greenhouse.io:8443/stripe",
            "https://jobs.lever.co@evil.example/metabase",
        ):
            with self.subTest(unsafe=unsafe), self.assertRaises(ValueError):
                greenhouse_token(unsafe)
        with self.assertRaises(ValueError):
            lever_slug("http://jobs.lever.co/metabase")

    def test_source_html_is_converted_to_inert_text_without_script_or_style_contents(self):
        text = clean_text('<p>Build <b>Python</b> APIs &amp; tools.</p><script>alert("xss")</script><style>.hidden{display:none}</style><img src=x onerror=alert(1)>')
        self.assertEqual(text, "Build Python APIs & tools.")
        self.assertNotIn("alert", text)
        self.assertNotIn("hidden", text)

    def test_source_links_keep_only_public_https_targets(self):
        self.assertEqual(safe_http_url("https://careers.example.com/jobs/7"), "https://careers.example.com/jobs/7")
        for unsafe in ("javascript:alert(1)", "http://example.com/job", "https://127.0.0.1/job", "https://localhost/job", "https://user:pass@example.com/job"):
            with self.subTest(unsafe=unsafe):
                self.assertEqual(safe_http_url(unsafe), "")

    def test_public_json_reader_rejects_redirects_and_oversized_bodies(self):
        requested: list[str] = []

        def redirect(request: httpx.Request) -> httpx.Response:
            requested.append(str(request.url))
            return httpx.Response(302, headers={"location": "http://127.0.0.1/internal"}, request=request)

        with self.assertRaises(httpx.HTTPStatusError):
            asyncio.run(fetch_public_json("https://boards-api.greenhouse.io/v1/boards/stripe/jobs", transport=httpx.MockTransport(redirect)))
        self.assertEqual(requested, ["https://boards-api.greenhouse.io/v1/boards/stripe/jobs"])

        def oversized(request: httpx.Request) -> httpx.Response:
            return httpx.Response(200, headers={"content-type": "application/json"}, content=b" " * (MAX_PUBLIC_RESPONSE_BYTES + 1), request=request)

        with self.assertRaisesRegex(ValueError, "too large"):
            asyncio.run(fetch_public_json("https://api.lever.co/v0/postings/example", transport=httpx.MockTransport(oversized)))

    def test_bounded_reader_accepts_the_current_arbeitnow_feed_size(self):
        # The official Arbeitnow feed is about 2.8 MB in production. It must
        # remain usable while the reader still enforces its explicit 4 MB cap.
        current_feed_sized_json = b"[]" + (b" " * (2_800_000 - 2))

        def current_arbeitnow_size(request: httpx.Request) -> httpx.Response:
            return httpx.Response(200, headers={"content-type": "application/json"}, content=current_feed_sized_json, request=request)

        result = asyncio.run(fetch_public_json("https://www.arbeitnow.com/api/job-board-api", transport=httpx.MockTransport(current_arbeitnow_size)))
        self.assertEqual(result, [])
        self.assertLess(len(current_feed_sized_json), MAX_PUBLIC_RESPONSE_BYTES)

    def test_greenhouse_reads_compact_index_then_only_bounded_detail_records(self):
        calls: list[str] = []

        async def provider_response(url: str, **_kwargs):
            calls.append(url)
            if "content=false" in url:
                return {"jobs": [{
                    "id": index, "title": f"Role {index}", "location": {"name": "Remote"},
                    "absolute_url": f"https://boards.greenhouse.io/example/jobs/{index}", "updated_at": "2026-01-01",
                } for index in range(20)]}
            job_id = url.rsplit("/", 1)[-1]
            return {
                "id": int(job_id), "title": f"Role {job_id}", "location": {"name": "Remote"},
                "absolute_url": f"https://boards.greenhouse.io/example/jobs/{job_id}",
                "content": f"<p>Build Python services for role {job_id}.</p><script>ignored()</script>", "updated_at": "2026-01-01",
            }

        with patch("app.sources.fetch_public_json", new=AsyncMock(side_effect=provider_response)):
            jobs = asyncio.run(fetch_greenhouse("example", limit=99))
        self.assertEqual(len(jobs), 12)
        self.assertIn("content=false", calls[0])
        self.assertEqual(len(calls), 13)  # one compact index plus 12 detail records
        self.assertTrue(all("ignored" not in job["description"] for job in jobs))

    def test_remotive_cache_filters_each_request_and_never_leaks_other_provider_jobs(self):
        remotive_job = {
            "id": "remotive-python", "origin": "public_feed", "source": "Remotive",
            "title": "Python Engineer", "company": "Remote Co", "description": "Build Python services.",
            "must_have": ["Python"], "preferred": [], "url": "https://example.com/remotive",
        }
        arbeitnow_job = {
            "id": "arbeitnow-python", "origin": "public_feed", "source": "Arbeitnow",
            "title": "Python Developer", "company": "Berlin Co", "description": "Build Python services.",
            "must_have": ["Python"], "preferred": [], "url": "https://example.com/arbeitnow",
        }
        main_module.REMOTIVE_CACHE.clear()
        main_module.REMOTIVE_CACHE_AT = None
        main_module.REMOTIVE_CACHE_ERROR = None
        try:
            with patch.object(main_module, "fetch_remotive", new=AsyncMock(return_value=[remotive_job])) as remotive_fetch, \
                 patch.object(main_module, "fetch_arbeitnow", new=AsyncMock(return_value=[arbeitnow_job])):
                remotive = asyncio.run(main_module.public_jobs(query="python", source="remotive"))
                self.assertEqual([item["id"] for item in remotive["jobs"]], ["remotive-python"])
                # A second, different query uses local filtering of the same raw cache.
                no_match = asyncio.run(main_module.public_jobs(query="golang", source="remotive"))
                self.assertEqual(no_match["jobs"], [])
                self.assertEqual(remotive_fetch.await_count, 1)
                arbeitnow = asyncio.run(main_module.public_jobs(query="python", source="arbeitnow"))
                self.assertEqual([item["id"] for item in arbeitnow["jobs"]], ["arbeitnow-python"])
        finally:
            main_module.REMOTIVE_CACHE.clear()
            main_module.REMOTIVE_CACHE_AT = None
            main_module.REMOTIVE_CACHE_ERROR = None

    def test_arbeitnow_query_cache_and_public_board_cache_are_bounded_and_summary_only(self):
        arbeitnow_job = {
            "id": "arbeitnow-cache-python", "origin": "public_cache", "source": "Arbeitnow", "title": "Python Developer",
            "company": "Example", "location": "Remote", "type": "Full time", "description": "Build Python services.",
            "must_have": ["Python"], "preferred": ["Docker"], "url": "https://careers.example.com/python",
        }
        board_job = {**arbeitnow_job, "id": "greenhouse-cache-python", "source": "Greenhouse"}
        main_module.ARBEITNOW_QUERY_CACHE.clear()
        main_module.PUBLIC_BOARD_CACHE.clear()
        try:
            with patch.object(main_module, "fetch_arbeitnow", new=AsyncMock(return_value=[arbeitnow_job])) as arbeitnow_fetch, \
                 patch.object(main_module, "fetch_greenhouse", new=AsyncMock(return_value=[board_job])) as greenhouse_fetch:
                first = asyncio.run(main_module.public_jobs(query="python", source="arbeitnow"))
                second = asyncio.run(main_module.public_jobs(query="python", source="arbeitnow"))
                self.assertEqual(arbeitnow_fetch.await_count, 1)
                self.assertEqual(first["provenance"]["cache_state"], "fresh")
                self.assertEqual(second["provenance"]["cache_state"], "cached")
                self.assertNotIn("description", second["jobs"][0])

                client = TestClient(app)
                board_first = client.post("/api/jobs/public-board", json={"source": "greenhouse", "board": "example"})
                board_second = client.post("/api/jobs/public-board", json={"source": "greenhouse", "board": "example"})
                self.assertEqual(board_first.status_code, 200, board_first.text)
                self.assertEqual(board_second.status_code, 200, board_second.text)
                self.assertEqual(greenhouse_fetch.await_count, 1)
                self.assertEqual(board_second.json()["provenance"]["cache_state"], "cached")
                self.assertNotIn("description", board_second.json()["jobs"][0])
        finally:
            main_module.ARBEITNOW_QUERY_CACHE.clear()
            main_module.PUBLIC_BOARD_CACHE.clear()
            main_module.JOBS.pop(arbeitnow_job["id"], None)
            main_module.JOBS.pop(board_job["id"], None)

    def test_independent_public_feeds_start_in_parallel_with_failure_isolation(self):
        remotive_job = {"id": "parallel-remotive", "origin": "public_cache", "source": "Remotive", "title": "Python role", "company": "Remote Co", "description": "Build Python.", "must_have": ["Python"], "preferred": [], "url": "https://example.com/remotive"}
        arbeitnow_job = {"id": "parallel-arbeitnow", "origin": "public_cache", "source": "Arbeitnow", "title": "Python role", "company": "Remote Co", "description": "Build Python.", "must_have": ["Python"], "preferred": [], "url": "https://example.com/arbeitnow"}
        main_module.REMOTIVE_CACHE.clear()
        main_module.REMOTIVE_CACHE_AT = None
        main_module.ARBEITNOW_QUERY_CACHE.clear()
        started: set[str] = set()
        release = asyncio.Event()

        async def delayed_remotive(**_kwargs):
            started.add("remotive")
            await release.wait()
            return [remotive_job]

        async def delayed_arbeitnow(**_kwargs):
            started.add("arbeitnow")
            await release.wait()
            return [arbeitnow_job]

        async def exercise():
            with patch.object(main_module, "fetch_remotive", new=AsyncMock(side_effect=delayed_remotive)), \
                 patch.object(main_module, "fetch_arbeitnow", new=AsyncMock(side_effect=delayed_arbeitnow)):
                task = asyncio.create_task(main_module.public_jobs(query="python", source="all"))
                await asyncio.sleep(0.03)
                both_started = started == {"remotive", "arbeitnow"}
                release.set()
                payload = await task
                return both_started, payload

        try:
            both_started, payload = asyncio.run(exercise())
            self.assertTrue(both_started, "both public feeds should start before either provider resolves")
            self.assertEqual({job["id"] for job in payload["jobs"]}, {"parallel-remotive", "parallel-arbeitnow"})
        finally:
            main_module.REMOTIVE_CACHE.clear()
            main_module.REMOTIVE_CACHE_AT = None
            main_module.ARBEITNOW_QUERY_CACHE.clear()
            main_module.JOBS.pop(remotive_job["id"], None)
            main_module.JOBS.pop(arbeitnow_job["id"], None)

    def test_live_queries_require_two_characters_on_both_route_contracts(self):
        client = TestClient(app)
        for path in ("/api/jobs/live?query=x", "/api/jobs?mode=live&query=%20"):
            with self.subTest(path=path):
                response = client.get(path)
                self.assertEqual(response.status_code, 422, response.text)
                self.assertIn("at least 2", response.json()["detail"])
        long_query = "x" * 121
        response = client.get(f"/api/jobs/live?query={long_query}")
        self.assertEqual(response.status_code, 422, response.text)
        self.assertIn("120 characters", response.json()["detail"])

    def test_static_deep_routes_are_shareable_and_unknown_browser_route_is_html_404(self):
        previous_static_dir = main_module.STATIC_DIR
        with tempfile.TemporaryDirectory() as directory:
            static_dir = Path(directory)
            (static_dir / "404.html").write_text("<html><title>HireSwarm — Not found</title><main>Not found</main></html>")
            for route in ("workspace", "applications", "practice", "documents"):
                (static_dir / f"{route}.html").write_text(f"<html><body>{route} route</body></html>")
            main_module.STATIC_DIR = static_dir
            try:
                client = TestClient(app)
                for route in ("workspace", "applications", "practice", "documents"):
                    response = client.get(f"/{route}")
                    self.assertEqual(response.status_code, 200, response.text)
                    self.assertIn(f"{route} route", response.text)
                missing = client.get("/not-a-real-page")
                self.assertEqual(missing.status_code, 404)
                self.assertTrue(missing.headers["content-type"].startswith("text/html"))
                self.assertIn("HireSwarm", missing.text)
                self.assertEqual(client.get("/api/not-a-route").status_code, 404)
                self.assertTrue(client.get("/api/not-a-route").headers["content-type"].startswith("application/json"))
                self.assertEqual(client.get("/missing.js").status_code, 404)
                self.assertTrue(client.get("/missing.js").headers["content-type"].startswith("application/json"))
            finally:
                main_module.STATIC_DIR = previous_static_dir

    def test_large_json_payloads_are_gzipped_for_slow_connections(self):
        response = TestClient(app).get("/api/jobs?mode=demo", headers={"Accept-Encoding": "gzip"})
        self.assertEqual(response.status_code, 200, response.text)
        # httpx transparently decodes the body, while the header confirms the
        # production middleware sent a compressed representation over the wire.
        self.assertEqual(response.headers.get("content-encoding"), "gzip")
        self.assertIn("Accept-Encoding", response.headers.get("vary", ""))
        self.assertGreaterEqual(len(response.json()["jobs"]), 3)

    def test_exported_static_assets_are_gzipped_and_immutable(self):
        previous_static_dir = main_module.STATIC_DIR
        with tempfile.TemporaryDirectory() as directory:
            static_dir = Path(directory)
            asset = static_dir / "_next" / "static" / "chunks" / "application.js"
            asset.parent.mkdir(parents=True)
            asset.write_text("const evidence = 'source-linked work example';\n" * 120)
            main_module.STATIC_DIR = static_dir
            try:
                response = TestClient(app).get("/_next/static/chunks/application.js", headers={"Accept-Encoding": "gzip"})
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.headers.get("content-encoding"), "gzip")
                self.assertIn("Accept-Encoding", response.headers.get("vary", ""))
                self.assertEqual(response.headers.get("cache-control"), "public, max-age=31536000, immutable")
            finally:
                main_module.STATIC_DIR = previous_static_dir

    def test_public_job_list_is_capped_and_long_description_stays_server_side_until_detail(self):
        raw_description = "x" * 9_000
        jobs = [{
            "id": f"public-{index}", "origin": "public_cache", "source": "Test source", "title": f"Role {index}",
            "company": "Example", "description": raw_description, "must_have": ["Python"], "preferred": [],
            "url": "https://careers.example.com/job", "retrieved_at": "2026-01-01T00:00:00+00:00",
        } for index in range(20)]
        main_module.REMOTIVE_CACHE.clear()
        main_module.REMOTIVE_CACHE_AT = None
        main_module.REMOTIVE_CACHE_ERROR = None
        main_module.ARBEITNOW_QUERY_CACHE.clear()
        try:
            with patch.object(main_module, "fetch_remotive", new=AsyncMock(return_value=jobs)), patch.object(main_module, "fetch_arbeitnow", new=AsyncMock(return_value=[])):
                payload = asyncio.run(main_module.public_jobs(query="role", source="all"))
            self.assertEqual(len(payload["jobs"]), main_module.PUBLIC_RESULT_LIMIT)
            self.assertTrue(all("description" not in item and "must_have" not in item and "preferred" not in item for item in payload["jobs"]))
            self.assertTrue(all(item["detail_loaded"] is False and item["skills"] == ["Python"] for item in payload["jobs"]))
            self.assertEqual(main_module.JOBS["public-0"]["description"], raw_description)
            detail = TestClient(app).get("/api/jobs/public-0")
            self.assertEqual(detail.status_code, 200, detail.text)
            self.assertTrue(detail.json()["detail_loaded"])
            self.assertEqual(len(detail.json()["description"]), main_module.CLIENT_DESCRIPTION_LIMIT)
        finally:
            main_module.REMOTIVE_CACHE.clear()
            main_module.REMOTIVE_CACHE_AT = None
            main_module.REMOTIVE_CACHE_ERROR = None
            main_module.ARBEITNOW_QUERY_CACHE.clear()
            for job in jobs:
                main_module.JOBS.pop(job["id"], None)

    def test_incomplete_applicant_is_not_silently_replaced_by_practice_evidence(self):
        applicant = CandidateInput(name="Real Applicant", resume_text="", evidence=[])
        self.assertEqual(candidate_as_dict(applicant)["name"], "Real Applicant")
        self.assertEqual(candidate_as_dict(applicant)["evidence"], [])

    def test_public_epoch_timestamp_is_human_readable(self):
        self.assertEqual(display_posted(0), "0")
        self.assertRegex(display_posted(1_700_000_000_000), r"^20\d\d-\d\d-\d\d$")

    def test_imported_text_yields_reviewable_evidence(self):
        candidate = _candidate_from_text(
            """Jane Applicant
Backend developer
EXPERIENCE
Built Python REST APIs using FastAPI and PostgreSQL for a student platform.
Dockerized the service and documented setup for three contributors.
Reduced manual reporting time by 40 percent with a React dashboard.
""",
            "jane.txt",
        )
        self.assertEqual(candidate.name, "Jane Applicant")
        self.assertGreaterEqual(len(candidate.evidence), 2)
        self.assertTrue(all(item.source_text for item in candidate.evidence))
        self.assertNotIn("Backend developer", [item.source_text for item in candidate.evidence])
        imported_skills = {skill for item in candidate.evidence for skill in item.skills}
        self.assertIn("FastAPI", imported_skills)
        self.assertIn("REST APIs", imported_skills)

    def test_cover_letter_uses_only_selected_evidence_and_names_gaps(self):
        candidate = {
            "name": "Taylor Example",
            "evidence": [{
                "evidence_id": "cv_01", "source_section": "Candidate-provided resume",
                "source_text": "Built a React dashboard for a campus operations team.",
                "skills": ["React"], "status": "verified",
            }],
        }
        job = {"title": "Backend Engineer", "company": "Example Co", "must_have": ["Python", "FastAPI"], "preferred": []}
        match = score_job(job, candidate)
        letter = build_cover_letter(job, candidate, match)
        ids = cover_letter_evidence_ids(job, candidate, match)
        self.assertIn("Built a React dashboard for a campus operations team.", letter)
        self.assertIn("do not have direct, verified evidence for Python, FastAPI", letter)
        self.assertNotIn("JWT authentication", letter)
        self.assertEqual(ids, ["cv_01"])

    def test_cv_upload_accepts_real_text_pdf_and_docx(self):
        resume = """Jane Applicant
Backend developer
EXPERIENCE
Built Python REST APIs using FastAPI and PostgreSQL for a student platform.
Dockerized the service and documented setup for three contributors.
Reduced manual reporting time by 40 percent with a React dashboard.
"""
        docx_buffer = BytesIO()
        docx = Document()
        for line in resume.splitlines():
            docx.add_paragraph(line)
        docx.save(docx_buffer)

        pdf_buffer = BytesIO()
        canvas = Canvas(pdf_buffer)
        y = 790
        for line in resume.splitlines():
            canvas.drawString(45, y, line)
            y -= 18
        canvas.save()

        client = TestClient(app)
        for filename, content, media_type in [
            ("jane.docx", docx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            ("jane.pdf", pdf_buffer.getvalue(), "application/pdf"),
        ]:
            with self.subTest(filename=filename):
                response = client.post("/api/candidate/upload", files={"file": (filename, content, media_type)})
                self.assertEqual(response.status_code, 200, response.text)
                payload = response.json()
                self.assertEqual(payload["import"]["format"], filename.rsplit(".", 1)[-1])
                self.assertGreaterEqual(payload["import"]["evidence_count"], 2)

    def test_unverified_note_does_not_become_application_proof(self):
        candidate = {"evidence": [{
            "evidence_id": "draft_01", "source_section": "Draft", "source_text": "Claimed FastAPI work.",
            "skills": ["FastAPI"], "status": "unverified",
        }]}
        match = score_job({"must_have": ["FastAPI"], "preferred": []}, candidate)
        self.assertEqual(match["score"], 0)
        self.assertEqual(match["verified_strengths"], [])
        self.assertEqual(build_safe_patches({"must_have": ["FastAPI"]}, candidate, match), [])
        self.assertFalse(validate_patches([{"evidence_ids": ["draft_01"]}], candidate)["safe_patches"])

    def test_no_matching_evidence_has_zero_fit_not_a_synthetic_floor(self):
        match = score_job({"must_have": ["Kubernetes"], "preferred": []}, {"evidence": []})
        self.assertEqual(match["score"], 0)
        self.assertEqual(match["coverage"], 0)

    def test_unmatched_role_ends_in_needs_evidence_not_export_ready(self):
        candidate = {"name": "Taylor Example", "headline": "Frontend", "location": "Remote", "evidence": [{
            "evidence_id": "cv_01", "source_section": "Experience", "source_text": "Built a React dashboard for a campus team.",
            "skills": ["React"], "status": "verified",
        }]}
        job = {"id": "qa-cloud", "title": "Cloud Engineer", "company": "QA Co", "must_have": ["Kubernetes", "Terraform"], "preferred": []}
        run = RunStatus(id="run-test", status="queued", selected_job_id="qa-cloud", mode="evidence_lab")
        with patch("app.main.asyncio.sleep", new=AsyncMock()):
            asyncio.run(execute_evidence_lab(run, job, candidate))
        self.assertEqual(run.status, "needs_evidence")
        self.assertEqual(run.result["patches"], [])
        self.assertEqual(run.events[-1].type, "evidence_needed")

    def test_same_evidence_statement_is_not_exported_twice(self):
        candidate = {
            "name": "Taylor Example",
            "evidence": [
                {"evidence_id": "cv_01", "source_section": "Experience", "source_text": "Developed REST APIs using Python and FastAPI for a student platform.", "skills": ["Python", "FastAPI", "REST APIs"], "status": "verified"},
                {"evidence_id": "cv_02", "source_section": "Project", "source_text": "Dockerized the service and documented local setup.", "skills": ["Docker"], "status": "verified"},
            ],
        }
        job = {"title": "Backend Engineer", "must_have": ["Python", "FastAPI", "REST APIs", "Docker"], "preferred": []}
        patches = build_safe_patches(job, candidate, score_job(job, candidate))
        self.assertEqual(len({patch["proposed"] for patch in patches}), len(patches))
        self.assertEqual(len(patches), 2)

    def test_target_requirement_order_controls_document_priority(self):
        candidate = copy.deepcopy(DEMO_CANDIDATE)
        job = copy.deepcopy(DEMO_JOBS[1])  # Full-stack role begins with React.
        patches = build_safe_patches(job, candidate, score_job(job, candidate))
        self.assertTrue(patches)
        self.assertIn("React", patches[0]["covered_requirements"])
        self.assertIn("React", patches[0]["proposed"])

    def test_non_backend_direct_requirement_receives_a_literal_patch(self):
        candidate = {"evidence": [{
            "evidence_id": "cv_01", "source_section": "Experience", "source_text": "Built a React dashboard for a campus operations team.",
            "skills": ["React"], "status": "verified",
        }]}
        job = {"title": "Frontend Engineer", "must_have": ["React"], "preferred": []}
        patches = build_safe_patches(job, candidate, score_job(job, candidate))
        self.assertEqual(len(patches), 1)
        self.assertEqual(patches[0]["proposed"], "Built a React dashboard for a campus operations team.")
        self.assertEqual(patches[0]["covered_requirements"], ["React"])

    def test_unrelated_evidence_is_not_used_to_pad_a_targeted_document(self):
        candidate = {"evidence": [{
            "evidence_id": "cv_01", "source_section": "Project", "source_text": "Dockerized a local development environment.",
            "skills": ["Docker"], "status": "verified",
        }]}
        job = {"title": "Platform Engineer", "must_have": ["Kubernetes", "Terraform"], "preferred": []}
        self.assertEqual(build_safe_patches(job, candidate, score_job(job, candidate)), [])

    def test_document_layer_defensively_removes_duplicate_bullets(self):
        result = self._result()
        result["patches"].append(copy.deepcopy(result["patches"][0]))
        self.assertEqual(len(_application_bullets(result)), len(result["patches"]) - 1)
        empty = {**result, "patches": []}
        self.assertFalse(export_readiness(empty)["passed"])

    def test_one_literal_can_cover_multiple_requirements_once(self):
        candidate = {"evidence": [{
            "evidence_id": "cv_01", "source_section": "Experience",
            "source_text": "Developed Dockerized REST APIs using Python and FastAPI for a student platform.",
            "skills": ["Python", "FastAPI", "REST APIs", "Docker"], "status": "verified",
        }]}
        job = {"title": "Backend Engineer", "must_have": ["Python", "FastAPI", "REST APIs", "Docker"], "preferred": []}
        patches = build_safe_patches(job, candidate, score_job(job, candidate))
        self.assertEqual(len(patches), 1)
        self.assertEqual(set(patches[0]["covered_requirements"]), {"Python", "FastAPI", "REST APIs", "Docker"})
        self.assertEqual(patches[0]["evidence_ids"], ["cv_01"])

    def test_normalize_rejects_empty_resume_instead_of_retaining_or_erasing_stale_evidence(self):
        client = TestClient(app)
        response = client.post("/api/candidate/normalize", json={
            "name": "Taylor Example", "headline": "Applicant", "location": "Remote", "preferences": [],
            "resume_text": "", 
            "evidence": [{"evidence_id": "old_01", "source_section": "Old file", "source_text": "Built a hidden stale project.", "skills": ["Python"], "status": "verified"}],
        })
        self.assertEqual(response.status_code, 422, response.text)
        self.assertIn("resume text", response.json()["detail"].lower())

    def test_normalize_rejects_text_without_reviewable_work_statement(self):
        response = TestClient(app).post("/api/candidate/normalize", json={
            "name": "Taylor Example", "headline": "Applicant", "location": "Remote", "preferences": [],
            "resume_text": "Taylor Example\nMotivated learner seeking an opportunity.", "evidence": [],
        })
        self.assertEqual(response.status_code, 422, response.text)
        self.assertIn("reviewable evidence", response.json()["detail"].lower())

    def test_empty_ledger_cannot_start_an_application_packet(self):
        client = TestClient(app)
        response = client.post("/api/runs", json={
            "job_id": "atlas-ai-backend",
            "candidate": {"name": "Taylor Example", "headline": "Applicant", "location": "Remote", "preferences": [], "resume_text": "", "evidence": []},
        })
        self.assertEqual(response.status_code, 422)
        self.assertIn("empty evidence ledger", response.json()["detail"])

    def test_read_only_match_preview_returns_honest_fit_without_creating_a_run(self):
        client = TestClient(app)
        candidate = copy.deepcopy(DEMO_CANDIDATE)
        before = set(main_module.RUNS)
        response = client.post("/api/match", json={"job_id": "atlas-ai-backend", "candidate": candidate})
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertIn("coverage", payload)
        self.assertIn("Docker", {item["skill"] for item in payload["verified_strengths"]})
        self.assertNotIn("Docker", {item["skill"] for item in payload["gaps"]})
        self.assertEqual(set(main_module.RUNS), before, "a fit preview must not silently start an application run")

    def test_match_preview_requires_reviewable_source_evidence(self):
        response = TestClient(app).post("/api/match", json={
            "job_id": "atlas-ai-backend",
            "candidate": {"name": "Taylor Example", "headline": "Applicant", "location": "Remote", "preferences": [], "resume_text": "", "evidence": []},
        })
        self.assertEqual(response.status_code, 422)
        self.assertIn("source-linked", response.json()["detail"])

    def test_run_requires_name_and_verified_evidence(self):
        client = TestClient(app)
        base = {"job_id": "atlas-ai-backend", "candidate": {"name": "", "headline": "Applicant", "location": "Remote", "preferences": [], "resume_text": "", "evidence": [{"evidence_id": "cv_01", "source_section": "Resume", "source_text": "Built a FastAPI service.", "skills": ["FastAPI"], "status": "verified"}]}}
        response = client.post("/api/runs", json=base)
        self.assertEqual(response.status_code, 422)
        self.assertIn("name", response.json()["detail"])
        base["candidate"]["name"] = "Taylor Example"
        base["candidate"]["evidence"][0]["status"] = "unverified"
        response = client.post("/api/runs", json=base)
        self.assertEqual(response.status_code, 422)
        self.assertIn("verified", response.json()["detail"])

    def test_manual_role_rejects_unsafe_listing_url(self):
        client = TestClient(app)
        response = client.post("/api/jobs/manual", json={
            "title": "Engineer", "company": "Example", "description": "Python work", "url": "javascript:alert(1)",
        })
        self.assertEqual(response.status_code, 422)
        self.assertIn("https://", response.json()["detail"])

    def test_manual_role_requires_public_https_url_and_never_fetches_it(self):
        client = TestClient(app)
        base = {"title": "Engineer", "company": "Example", "description": "Build Python services with clear ownership.", "url": ""}
        for unsafe in ("http://careers.example.com/job", "https://localhost/job", "https://127.0.0.1/job", "https://user:pass@careers.example.com/job"):
            with self.subTest(unsafe=unsafe):
                response = client.post("/api/jobs/manual", json={**base, "url": unsafe})
                self.assertEqual(response.status_code, 422, response.text)
        response = client.post("/api/jobs/manual", json={**base, "url": "https://careers.example.com/job"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["url"], "https://careers.example.com/job")

    def test_manual_role_extracts_extended_controlled_skill_taxonomy(self):
        client = TestClient(app)
        response = client.post("/api/jobs/manual", json={
            "title": "Platform Engineer", "company": "Example Co",
            "description": "Build Terraform modules and Apache Airflow workflows, write testing plans, and ship LLM applications with HTML and CSS interfaces.",
        })
        self.assertEqual(response.status_code, 200, response.text)
        extracted = response.json()["must_have"] + response.json()["preferred"]
        for skill in ["Terraform", "Airflow", "Testing", "LLM applications", "HTML", "CSS"]:
            self.assertIn(skill, extracted)

    def test_approval_rechecks_document_qa_and_rejects_zero_patch_packet(self):
        run_id = "run-malformed-approval-test"
        main_module.RUNS[run_id] = RunStatus(
            id=run_id, status="awaiting_approval", selected_job_id="atlas-ai-backend", mode="evidence_lab",
            result={
                "candidate": {"name": "Taylor Example"},
                "job": {"title": "Engineer", "company": "Example"},
                "patches": [],
                "cover_letter": "A draft.",
                # Simulate an untrusted stale flag; server-side approval must ignore it.
                "export_readiness": {"passed": True},
            },
        )
        try:
            response = TestClient(app).post(f"/api/runs/{run_id}/approve")
            self.assertEqual(response.status_code, 409, response.text)
            self.assertIn("Document QA", response.json()["detail"])
            self.assertEqual(main_module.RUNS[run_id].status, "awaiting_approval")
        finally:
            main_module.RUNS.pop(run_id, None)

    def test_approval_rejects_a_patch_with_forged_evidence_id(self):
        run_id = "run-forged-evidence-test"
        result = self._result()
        result["patches"][0]["evidence_ids"] = ["not-in-the-candidate-ledger"]
        main_module.RUNS[run_id] = RunStatus(
            id=run_id, status="awaiting_approval", selected_job_id="atlas-ai-backend", mode="evidence_lab", result=result,
        )
        try:
            response = TestClient(app).post(f"/api/runs/{run_id}/approve")
            self.assertEqual(response.status_code, 409, response.text)
            readiness = main_module.RUNS[run_id].result["export_readiness"]
            integrity = next(item for item in readiness["checks"] if item["label"] == "Server-side evidence integrity")
            self.assertFalse(integrity["passed"])
        finally:
            main_module.RUNS.pop(run_id, None)

    def test_approval_is_idempotent_for_double_click_tabs_refresh_and_sse(self):
        run_id = "run-idempotent-approval-test"
        main_module.RUNS[run_id] = RunStatus(
            id=run_id, status="awaiting_approval", selected_job_id="atlas-ai-backend", mode="evidence_lab", result=self._result(),
        )
        try:
            def approve_once():
                with TestClient(app) as client:
                    return client.post(f"/api/runs/{run_id}/approve")

            with ThreadPoolExecutor(max_workers=2) as pool:
                responses = list(pool.map(lambda _: approve_once(), range(2)))
            self.assertTrue(all(response.status_code == 200 for response in responses), [response.text for response in responses])
            payloads = [response.json() for response in responses]
            self.assertEqual(sum(bool(payload.get("already_approved")) for payload in payloads), 1)
            self.assertEqual(sum(event.type == "approved" for event in main_module.RUNS[run_id].events), 1)

            # A page refresh / repeated click receives the same approved state,
            # and the terminal SSE history contains only the single approval.
            with TestClient(app) as client:
                refreshed = client.post(f"/api/runs/{run_id}/approve")
                checkpoint = client.get(f"/api/runs/{run_id}")
                stream = client.get(f"/api/runs/{run_id}/events")
            self.assertEqual(refreshed.status_code, 200)
            self.assertTrue(refreshed.json()["already_approved"])
            self.assertEqual(checkpoint.status_code, 200)
            self.assertEqual(checkpoint.json()["status"], "approved")
            self.assertEqual(checkpoint.json()["event_count"], 1)
            self.assertEqual(stream.status_code, 200)
            self.assertEqual(stream.text.count('"type": "approved"'), 1)
        finally:
            main_module.RUNS.pop(run_id, None)

    def test_approved_export_formats_are_real_and_pdf_round_trips_to_text(self):
        result = self._result()
        pdf = build_pdf(result)
        self.assertTrue(pdf.startswith(b"%PDF"))
        extracted = " ".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf)).pages)
        self.assertIn("Hussain Ahmed", extracted)
        self.assertIn("TARGETED EXPERIENCE", extracted)

        docx = build_docx(result)
        document = Document(BytesIO(docx))
        self.assertIn("Hussain Ahmed", "\n".join(paragraph.text for paragraph in document.paragraphs))

        readiness = export_readiness(result)
        self.assertTrue(readiness["passed"])
        self.assertTrue(next(item for item in readiness["checks"] if item["label"] == "PDF text round-trip")["passed"])


if __name__ == "__main__":
    unittest.main()
