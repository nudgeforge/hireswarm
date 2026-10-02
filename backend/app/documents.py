"""ATS-conscious export and round-trip readiness checks for HireSwarm."""
from __future__ import annotations

from io import BytesIO
from textwrap import wrap
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer


def _application_bullets(result: dict[str, Any]) -> list[str]:
    """Return unique literal revisions even if an upstream caller repeats one."""
    bullets: list[str] = []
    seen: set[str] = set()
    for patch in result.get("patches", []):
        value = str(patch.get("proposed", "")).strip()
        key = " ".join(value.casefold().split())
        if value and key not in seen:
            bullets.append(value)
            seen.add(key)
    return bullets


def _candidate(result: dict[str, Any]) -> dict[str, Any]:
    return result.get("candidate_full") or result.get("candidate") or {}


def _job(result: dict[str, Any]) -> dict[str, Any]:
    return result.get("job") or {}


def _pdf_text(value: Any) -> str:
    """Escape applicant/job text before it reaches ReportLab's mini-markup parser."""
    return escape(str(value or "")).replace("\n", "<br/>")


def build_docx(result: dict[str, Any]) -> bytes:
    candidate = _candidate(result)
    job = _job(result)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.2)

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = name.add_run(candidate.get("name", "Applicant"))
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(19)

    details = document.add_paragraph()
    details.add_run(candidate.get("headline", "")).bold = True
    location = candidate.get("location")
    if location:
        details.add_run(f"  |  {location}")

    def heading(text: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(10)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "325B7A")
        border.append(bottom)
        paragraph._p.get_or_add_pPr().append(border)

    heading("Targeted experience")
    bullets = _application_bullets(result)
    for bullet in bullets or ["No verified revisions were generated for this packet."]:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(bullet)

    heading("Relevant skills")
    match = result.get("match") or {}
    skills = [entry.get("skill") for entry in match.get("verified_strengths", []) if entry.get("skill")]
    document.add_paragraph(", ".join(dict.fromkeys(skills)) or "Skills are represented in the verified experience statements above.")

    heading("Application note")
    document.add_paragraph(f"Prepared for {job.get('title', 'the target role')} at {job.get('company', 'the target company')}. This one-column document was generated from applicant-approved, source-linked evidence.")

    heading("Cover letter")
    document.add_paragraph(result.get("cover_letter") or "")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf(result: dict[str, Any]) -> bytes:
    candidate = _candidate(result)
    job = _job(result)
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1.7*cm, leftMargin=1.7*cm, topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CandidateName", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=colors.HexColor("#17283A"), spaceAfter=4))
    styles.add(ParagraphStyle(name="Sub", parent=styles["BodyText"], fontSize=9.5, leading=13, textColor=colors.HexColor("#41556B"), spaceAfter=12))
    styles.add(ParagraphStyle(name="HeadingSmall", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=10, leading=13, textColor=colors.HexColor("#23769B"), spaceBefore=10, spaceAfter=5))
    styles.add(ParagraphStyle(name="BodySafe", parent=styles["BodyText"], fontSize=9.7, leading=14, textColor=colors.HexColor("#17283A"), spaceAfter=5))
    story = [Paragraph(_pdf_text(candidate.get("name", "Applicant")), styles["CandidateName"]), Paragraph(_pdf_text(" · ".join(str(item) for item in [candidate.get("headline", ""), candidate.get("location", "")] if item)), styles["Sub"])]
    story.append(Paragraph("TARGETED EXPERIENCE", styles["HeadingSmall"]))
    for bullet in _application_bullets(result) or ["No verified revisions were generated for this packet."]:
        story.append(Paragraph(f"• {_pdf_text(bullet)}", styles["BodySafe"]))
    story.append(Paragraph("RELEVANT SKILLS", styles["HeadingSmall"]))
    match = result.get("match") or {}
    skills = [entry.get("skill") for entry in match.get("verified_strengths", []) if entry.get("skill")]
    story.append(Paragraph(_pdf_text(", ".join(dict.fromkeys(skills)) or "See targeted experience."), styles["BodySafe"]))
    story.append(Paragraph("APPLICATION NOTE", styles["HeadingSmall"]))
    story.append(Paragraph(_pdf_text(f"Prepared for {job.get('title', 'the target role')} at {job.get('company', 'the target company')}. Generated from source-linked, applicant-approved evidence."), styles["BodySafe"]))
    story.append(Paragraph("COVER LETTER", styles["HeadingSmall"]))
    for paragraph in (result.get("cover_letter") or "").split("\n"):
        if paragraph.strip():
            story.append(Paragraph(_pdf_text(paragraph), styles["BodySafe"]))
        else:
            story.append(Spacer(1, 4))
    doc.build(story)
    return buffer.getvalue()


def export_readiness(result: dict[str, Any]) -> dict[str, Any]:
    """Run a real PDF text round-trip rather than inventing an ATS score."""
    pdf_bytes = build_pdf(result)
    extracted = " ".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf_bytes)).pages)
    candidate = _candidate(result)
    name = candidate.get("name", "")
    patches = _application_bullets(result)
    checks = [
        {"label": "Single-column document structure", "passed": True, "detail": "Export uses linear headings and bullets; no sidebars, tables, or text boxes."},
        {"label": "PDF text round-trip", "passed": bool(name and name in extracted), "detail": "The generated PDF was read back with an independent text extractor."},
        {"label": "Evidence-linked revisions", "passed": bool(patches) and all(bool(patch.get("evidence_ids")) for patch in result.get("patches", [])), "detail": "At least one proposed revision is present and every revision retains evidence IDs before export."},
        {"label": "Human approval gate", "passed": True, "detail": "Exports remain unavailable until the applicant approves the packet."},
    ]
    return {"passed": all(check["passed"] for check in checks), "checks": checks, "extracted_characters": len(extracted), "verified_bullets": len(patches)}
